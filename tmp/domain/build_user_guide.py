from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from zipfile import ZipFile
import re
base=Path.cwd();out=base/'output/doc';out.mkdir(parents=True,exist_ok=True)
doc=Document()
s=doc.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7)
s.top_margin=s.bottom_margin=Cm(2);s.left_margin=s.right_margin=Cm(2.2)
normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(12)
for sn in ['Title','Heading 1','Heading 2','Heading 3']:
 doc.styles[sn].font.color.rgb=RGBColor(0,0,0)
def fmt(p,text,size=12,bold=False):
 p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 pp=p._p.get_or_add_pPr();b=OxmlElement('w:bidi');b.set(qn('w:val'),'1');pp.append(b)
 p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_after=Pt(6)
 p.paragraph_format.widow_control=True
 for part in re.split(r'([A-Za-z0-9][A-Za-z0-9 .&?]*[A-Za-z0-9?.]|[A-Za-z0-9])',text):
  if not part: continue
  r=p.add_run(part);r.font.name='Arial';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor(0,0,0);r.font.rtl=not bool(re.match('[A-Za-z0-9]',part))
  rp=r._r.get_or_add_rPr();rp.rFonts.set(qn('w:cs'),'Arial')
  z=OxmlElement('w:szCs');z.set(qn('w:val'),str(2*size));rp.append(z)
  z=OxmlElement('w:bCs');z.set(qn('w:val'),'1' if bold else '0');rp.append(z)
for line in (base/'tmp/domain/user_guide.txt').read_text(encoding='utf-8-sig').splitlines():
 if not line: continue
 level=len(line)-len(line.lstrip('#'))
 style={0:'Normal',1:'Title',2:'Heading 2',3:'Heading 3'}[level]
 p=doc.add_paragraph(style=style);fmt(p,line.lstrip('# ').strip(),{0:12,1:23,2:17,3:13}[level],level>0)
 if level:
  p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(14 if level!=1 else 0)
 if re.match(r'^\d+\.',line): p.paragraph_format.space_after=Pt(4)
fmt(s.header.paragraphs[0],'LifeTracker',10)
f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=f.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.addnext(field)
standalone=out/'LifeTracker_מדריך_למשתמש.docx';doc.save(standalone)
source=out/'LifeTracker_4_כולל_תיאור_תחום_הידע.docx';merged=out/'LifeTracker_5_כולל_מדריך_למשתמש.docx'
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with ZipFile(source) as z,ZipFile(merged,'w') as target:
 root=etree.fromstring(z.read('word/document.xml'));body=root.find('w:body',ns)
 insertion=len(body)-1
 breaker=OxmlElement('w:p');r=OxmlElement('w:r');br=OxmlElement('w:br');br.set(qn('w:type'),'page');r.append(br);breaker.append(r);body.insert(insertion,breaker);insertion+=1
 for i,p in enumerate(doc.paragraphs):
  elem=deepcopy(p._p)
  if i==0: elem.find('w:pPr/w:pStyle',ns).set(qn('w:val'),'Heading1')
  elif p.style.name=='Normal':
   st=elem.find('w:pPr/w:pStyle',ns)
   if st is not None:st.getparent().remove(st)
  body.insert(insertion,elem);insertion+=1
 for entry in z.infolist():
  target.writestr(entry,etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True) if entry.filename=='word/document.xml' else z.read(entry.filename))
old=Document(source);new=Document(merged)
assert [p.text for p in new.paragraphs[:len(old.paragraphs)]]==[p.text for p in old.paragraphs]
assert [p.text for p in new.paragraphs[len(old.paragraphs)+1:]]==[p.text for p in doc.paragraphs]
for path in [standalone,merged]:
 with ZipFile(path) as z: assert z.testzip() is None
 print(path)
print('Verified original content preserved, new chapter matches, and both files are valid.')
print('Guide word count:',sum(len(p.text.split()) for p in doc.paragraphs))
